"""Spreadsheet, Word and image adapters → the same Grid model used for PDFs."""
from __future__ import annotations

import csv
import io
import os
import subprocess
import tempfile
from datetime import datetime, time

from app.config import get_settings
from app.extraction.model import Finding, Grid, GridCell
from app.extraction.ocr import ocr_available, ocr_grid
from app.extraction.timeparse import parse_day


class ConversionError(Exception):
    pass


def convert_with_soffice(data: bytes, src_ext: str, target: str) -> bytes:
    """Convert legacy .xls/.doc with LibreOffice headless in an isolated temp dir."""
    s = get_settings()
    with tempfile.TemporaryDirectory(prefix="tt-convert-") as d:
        src = os.path.join(d, "input" + src_ext)
        with open(src, "wb") as fh:
            fh.write(data)
        try:
            subprocess.run(
                [s.soffice_path, "--headless", "--norestore", f"-env:UserInstallation=file://{d}/profile",
                 "--convert-to", target, "--outdir", d, src],
                check=True, capture_output=True, timeout=120,
            )
        except (subprocess.SubprocessError, FileNotFoundError) as e:
            raise ConversionError("The legacy Office file could not be converted.") from e
        out = os.path.join(d, "input." + target)
        if not os.path.exists(out):
            raise ConversionError("The legacy Office file could not be converted.")
        with open(out, "rb") as fh:
            return fh.read()


def _cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, time):
        return v.strftime("%H:%M")
    if isinstance(v, datetime):
        return v.strftime("%H:%M") if v.date().year < 1901 else v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _grids_from_matrix(matrix: list[list[str]], spans: dict[tuple[int, int], tuple[int, int]],
                       covered: set[tuple[int, int]], page: int, method: str) -> list[Grid]:
    """Find weekday header rows in a text matrix and cut out each timetable block."""
    grids: list[Grid] = []
    n = len(matrix)
    headers = [r for r in range(n) if len({parse_day(c) for c in matrix[r] if c and parse_day(c)}) >= 3]
    if not headers:
        # transposed layout: days down one column → transpose and retry once
        width = max((len(r) for r in matrix), default=0)
        for c in range(min(width, 3)):
            if len({parse_day(r[c]) for r in matrix if c < len(r) and r[c] and parse_day(r[c])}) >= 3:
                t = [[matrix[r][c] if c < len(matrix[r]) else "" for r in range(n)] for c in range(width)]
                tspans = {(c, r): (cs, rs) for (r, c), (rs, cs) in spans.items()}
                tcov = {(c, r) for (r, c) in covered}
                return _grids_from_matrix(t, tspans, tcov, page, method)
        return []
    for i, h in enumerate(headers):
        end = headers[i + 1] if i + 1 < len(headers) else n
        # stop the block at the first completely empty row after content
        last = h
        for r in range(h + 1, end):
            if any(matrix[r]):
                last = r
            elif last > h:
                break
        prev_end = headers[i - 1] + 1 if i else 0
        header_text = "\n".join(" ".join(c for c in matrix[r] if c) for r in range(prev_end if i == 0 else prev_end, h) if any(matrix[r]))
        after = [matrix[r] for r in range(last + 1, end) if any(matrix[r])]
        cells: list[GridCell] = []
        for r in range(h, last + 1):
            for c, val in enumerate(matrix[r]):
                if (r, c) in covered:
                    continue
                rs, cs = spans.get((r, c), (1, 1))
                rs = min(rs, last + 1 - r)
                if val or (r == h):
                    cells.append(GridCell(r - h, c, val, rs, cs))
        width = max((len(matrix[r]) for r in range(h, last + 1)), default=0)
        grids.append(Grid(page=page, method=method, n_rows=last - h + 1, n_cols=width, cells=cells,
                          header_text=header_text, context_text="\n".join(" ".join(c for c in row if c) for row in after),
                          legend_rows=[[c for c in row if c] for row in after]))
    return grids


def extract_xlsx(data: bytes) -> tuple[list[Grid], list[dict], list[Finding], int]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=False)
    grids, pages, findings = [], [], []
    for idx, ws in enumerate(wb.worksheets, start=1):
        rows = [[_cell_str(v) for v in row] for row in ws.iter_rows(values_only=True)]
        spans, covered = {}, set()
        for mr in ws.merged_cells.ranges:
            r0, c0 = mr.min_row - 1, mr.min_col - 1
            spans[(r0, c0)] = (mr.max_row - mr.min_row + 1, mr.max_col - mr.min_col + 1)
            for r in range(mr.min_row - 1, mr.max_row):
                for c in range(mr.min_col - 1, mr.max_col):
                    if (r, c) != (r0, c0):
                        covered.add((r, c))
        g = _grids_from_matrix(rows, spans, covered, idx, "SPREADSHEET")
        pages.append({"page": idx, "sheet": ws.title, "method": "SPREADSHEET", "status": "PARSED" if g else "NO_TABLE", "grids": len(g)})
        if not g:
            findings.append(Finding("PAGE_ERROR", "NO_TIMETABLE_GRID", f"Sheet '{ws.title}' has no weekday/time grid.", idx))
        grids += g
    return grids, pages, findings, len(wb.worksheets)


def extract_csv(data: bytes) -> tuple[list[Grid], list[dict], list[Finding], int]:
    text = data.decode("utf-8-sig", errors="replace")
    dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t") if text.strip() else csv.excel
    rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), dialect)]
    g = _grids_from_matrix(rows, {}, set(), 1, "SPREADSHEET")
    findings = [] if g else [Finding("PAGE_ERROR", "NO_TIMETABLE_GRID", "The CSV has no weekday/time grid.", 1)]
    return g, [{"page": 1, "method": "SPREADSHEET", "status": "PARSED" if g else "NO_TABLE", "grids": len(g)}], findings, 1


def extract_docx(data: bytes) -> tuple[list[Grid], list[dict], list[Finding], int]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(data))
    grids, pages, findings = [], [], []
    preceding: list[str] = []
    tables_seen = 0
    body = doc.element.body
    all_tables: list[tuple[Table, str]] = []
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            t = Paragraph(child, doc).text.strip()
            if t:
                preceding.append(t)
        elif tag == "tbl":
            all_tables.append((Table(child, doc), "\n".join(preceding[-8:])))
            preceding = []
    trailing = "\n".join(preceding)
    legend_rows: list[list[str]] = []
    tt_grids: list[Grid] = []
    for table, header in all_tables:
        tables_seen += 1
        rows = len(table.rows)
        cols = len(table.columns)
        matrix: list[list[str]] = []
        ids: list[list[object]] = []  # the lxml <w:tc> elements themselves (not id(): proxies get recycled)
        for r in range(rows):
            row_txt, row_ids = [], []
            for c in range(cols):
                try:
                    cell = table.cell(r, c)
                except IndexError:
                    row_txt.append("")
                    row_ids.append(None)
                    continue
                row_ids.append(cell._tc)
                row_txt.append(cell.text.strip())
            matrix.append(row_txt)
            ids.append(row_ids)
        spans, covered, seen = {}, set(), {}
        for r in range(rows):
            for c in range(cols):
                tc = ids[r][c]
                if tc is None:
                    continue
                if tc in seen:
                    r0, c0 = seen[tc]
                    covered.add((r, c))
                    rs, cs = spans.get((r0, c0), (1, 1))
                    spans[(r0, c0)] = (max(rs, r - r0 + 1), max(cs, c - c0 + 1))
                else:
                    seen[tc] = (r, c)
        for (r, c) in covered:
            matrix[r][c] = ""
        g = _grids_from_matrix(matrix, spans, covered, tables_seen, "DOCX")
        if g:
            g[0].header_text = (header + "\n" + g[0].header_text).strip()
            tt_grids += g
        else:
            legend_rows += [[x for x in row if x] for row in matrix]
    for g in tt_grids:
        g.legend_rows = g.legend_rows + legend_rows
        g.context_text = (g.context_text + "\n" + trailing).strip()
        pages.append({"page": g.page, "method": "DOCX", "status": "PARSED", "grids": 1})
    grids = tt_grids
    if not grids:
        findings.append(Finding("BLOCKING", "NO_TIMETABLE_GRID", "No table with a weekday header was found in the document."))
    return grids, pages, findings, max(tables_seen, 1)


def extract_image(data: bytes) -> tuple[list[Grid], list[dict], list[Finding], int]:
    if not ocr_available():
        return [], [{"page": 1, "method": "OCR", "status": "FAILED"}], [
            Finding("BLOCKING", "OCR_UNAVAILABLE", "Image uploads need OCR, which is not available on this server.")], 1
    g, words, extra = ocr_grid(data, 1)
    confs = [w.conf for w in words if w.conf is not None]
    info = {**extra, "page": 1, "method": "OCR", "status": "PARSED" if g else "FAILED", "grids": 1 if g else 0,
            "ocr_mean_confidence": round(sum(confs) / len(confs), 1) if confs else None}
    findings = [] if g else [Finding("PAGE_ERROR", "OCR_NO_TIMETABLE_FOUND",
                                     "OCR could not recognise a weekday/time grid in the image.", 1)]
    return ([g] if g else []), [info], findings, 1
