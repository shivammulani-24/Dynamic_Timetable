"""Generates SYNTHETIC timetable fixtures for regression tests.

These files imitate the *layout traits* described in the PDF Extraction Specification (multiple
sections per file, Tentative division, combined BE divisions, per-page legends, stacked lab
entries, merged two-hour labs, breaks, a garbled '12.15 a.m. to 01.15 a.m.' row label that cannot be resolved
without review). They are NOT the
college's real timetable; all names/codes are invented. The college's real PDF is a separate fixture
(tests/fixtures/reference/college_timetable.pdf).

Run:  python -m app.devdata.synthetic [out_dir]
"""
from __future__ import annotations

import csv
import io
import os

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(os.getcwd(), "var", "synthetic")

DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
TIMES = ["9.00 a.m. to 10.00 a.m.", "10.00 a.m. to 11.00 a.m.", "11.00 a.m. to 11.15 a.m.",
         "11.15 a.m. to 12.15 p.m.", "12.15 a.m. to 01.15 a.m.", "01.15 p.m. to 02.00 p.m.",
         "02.00 p.m. to 03.00 p.m.", "03.00 p.m. to 04.00 p.m."]
FACULTY = [("KKD", "Prof. Kiran K. Desai"), ("AVN", "Prof. Anita V. Nair"), ("PJB", "Dr. Prakash J. Bhatt"),
           ("RHS", "Prof. Rhea H. Shah")]
SUBJECTS = [("DBMS", "Database Management Systems"), ("COA", "Computer Organization and Architecture"),
            ("DS", "Data Structures"), ("DSGT", "Discrete Structures and Graph Theory")]

# rows x days; None = empty; "SPAN" = covered by the cell above (two-hour lab)
SE_A = [
    ["DBMS/KKD/508", "COA/AVN/509 old", "DS/PJB/508", "DSGT/RHS/702-B", "DBMS/KKD/508"],
    ["DS Lab A1/AVN/604\nDS Lab A2/PJB/605", "DBMS/KKD/508", "COA/AVN/509 old", "DS/PJB/508", None],
    ["SHORT BREAK"] * 5,
    ["COA/AVN/509 old", "DSGT/RHS/702-B", "DBMS Lab A3/KKD/606-4\nCOA Lab A4/AVN/607", "DBMS/KKD/508", "DS/PJB/508"],
    ["DSGT/RHS/702-B", None, "SPAN", "COA/AVN/509 old", "DSGT/RHS/702-B"],
    ["LONG BREAK"] * 5,
    ["Student Activity Slot", "DS/PJB/508", "DSGT/RHS/702-B", None, "COA/AVN/509 old"],
    ["Faculty Quality Circle Slot", "QQ/KKD", None, "DS/PJB/508", None],
]
SE_C = [
    ["COA/AVN/510", "DBMS/KKD/510", None, "DS/PJB/510", "DSGT/RHS/510"],
    ["DBMS/KKD/510", None, "COA/AVN/510", None, "DS/PJB/510"],
    ["SHORT BREAK"] * 5,
    ["DS/PJB/510", "COA/AVN/510", "DBMS/KKD/510", "DSGT/RHS/510", None],
    [None, "DSGT/RHS/510", "DS/PJB/510", "COA/AVN/510", "DBMS/KKD/510"],
    ["LONG BREAK"] * 5,
    ["DSGT/RHS/510", None, "DBMS/KKD/510", None, "COA/AVN/510"],
    [None, "DS/PJB/510", None, "DBMS/KKD/510", None],
]
BE = [
    ["Open Elective", "PE III-TSDA Batch1/RHS/701\nPE III-CC Batch2/PJB/702", None, "DBMS/KKD/703", None],
    ["DS Lab B2/AVN/604", None, "COA/AVN/703", None, "DS/PJB/703"],
    ["SHORT BREAK"] * 5,
    [None, "DSGT/RHS/703", None, "COA/AVN/703", None],
    ["DBMS/KKD/703", None, "DS/PJB/703", None, "DSGT/RHS/703"],
    ["LONG BREAK"] * 5,
    [None, "COA/AVN/703", None, "DS/PJB/703", None],
    [None, None, None, None, None],
]
PAGES = [
    ("SE BTech Computer Engineering (Division-A)", SE_A),
    ("SE BTech Computer Engineering (Division-C) Tentative", SE_C),
    ("BE BTech Computer Engineering (Division-A, B, C, D)", BE),
]
HEADER = ["Example Institute of Technology (SYNTHETIC TEST DATA)", "Department of Computer Engineering",
          "AY 2026-27  ODD Semester    W.E.F. 10 August 2026"]


def make_pdf(path: str) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    cell_style = styles["BodyText"].clone("cell", fontSize=7, leading=8.5)
    story = []
    for title, grid in PAGES:
        for h in HEADER:
            story.append(Paragraph(h, styles["Normal"]))
        story.append(Paragraph(title, styles["Heading3"]))
        data = [["Time"] + DAYS]
        spans = []
        for r, (t, row) in enumerate(zip(TIMES, grid), start=1):
            out = [t]
            for c, v in enumerate(row, start=1):
                if v == "SPAN":
                    out.append("")
                    spans.append(("SPAN", (c, r - 1), (c, r)))
                else:
                    out.append(Paragraph((v or "").replace("\n", "<br/>"), cell_style))
            data.append(out)
            if row[0] and "BREAK" in row[0]:
                spans.append(("SPAN", (1, r), (5, r)))
        tbl = Table(data, colWidths=[95] + [130] * 5, repeatRows=1)
        tbl.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black), ("FONTSIZE", (0, 0), (-1, -1), 7),
                                 ("VALIGN", (0, 0), (-1, -1), "MIDDLE")] + spans))
        story.append(tbl)
        story.append(Spacer(1, 6))
        story.append(Paragraph("Faculty", styles["Normal"]))
        story.append(Paragraph("    ".join(f"{a} - {b}" for a, b in FACULTY), cell_style))
        story.append(Paragraph("Subjects", styles["Normal"]))
        story.append(Paragraph("    ".join(f"{a} - {b}" for a, b in SUBJECTS), cell_style))
        story.append(PageBreak())
    SimpleDocTemplate(path, pagesize=landscape(A4), leftMargin=20, rightMargin=20, topMargin=20, bottomMargin=20).build(story)


def _matrix(title, grid):
    rows = [[h] for h in HEADER] + [[title], ["Time"] + DAYS]
    for t, row in zip(TIMES, grid):
        rows.append([t] + [("" if v in (None, "SPAN") else v) for v in row])
    return rows


def make_xlsx(path: str) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for i, (title, grid) in enumerate(PAGES[:2]):
        ws = wb.create_sheet(f"Sheet{i + 1}")
        for row in _matrix(title, grid):
            ws.append(row)
        base = len(HEADER) + 2  # header lines + title + day header (1-based next row)
        for r, row in enumerate(grid):
            for c, v in enumerate(row):
                if v == "SPAN":
                    ws.merge_cells(start_row=base + r, start_column=c + 2, end_row=base + r + 1, end_column=c + 2)
        ws.append([])
        ws.append(["Faculty"])
        for a, b in FACULTY:
            ws.append([a, b])
        ws.append(["Subjects"])
        for a, b in SUBJECTS:
            ws.append([a, b])
    wb.save(path)


def make_csv(path: str) -> None:
    title, grid = PAGES[0]
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh)
        for row in _matrix(title, grid):
            w.writerow(row)
        w.writerow([])
        for a, b in FACULTY + SUBJECTS:
            w.writerow([f"{a} - {b}"])


def make_docx(path: str) -> None:
    from docx import Document

    doc = Document()
    title, grid = PAGES[0]
    for h in HEADER + [title]:
        doc.add_paragraph(h)
    t = doc.add_table(rows=len(TIMES) + 1, cols=6)
    for c, d in enumerate(["Time"] + DAYS):
        t.cell(0, c).text = d
    for r, (tm, row) in enumerate(zip(TIMES, grid), start=1):
        t.cell(r, 0).text = tm
        for c, v in enumerate(row, start=1):
            if v and v != "SPAN":
                t.cell(r, c).text = v
    for r, row in enumerate(grid, start=1):
        for c, v in enumerate(row, start=1):
            if v == "SPAN":
                t.cell(r - 1, c).merge(t.cell(r, c))
    doc.add_paragraph("Faculty")
    for a, b in FACULTY:
        doc.add_paragraph(f"{a} - {b}")
    doc.add_paragraph("Subjects")
    for a, b in SUBJECTS:
        doc.add_paragraph(f"{a} - {b}")
    doc.save(path)


def make_image_and_scanned(pdf_path: str, png_path: str, scanned_path: str) -> None:
    import pymupdf as fitz

    src = fitz.open(pdf_path)
    pix = src[0].get_pixmap(dpi=200)
    pix.save(png_path)
    out = fitz.open()
    page = out.new_page(width=src[0].rect.width, height=src[0].rect.height)
    page.insert_image(page.rect, stream=pix.tobytes("jpeg", jpg_quality=85))
    out.save(scanned_path)


def make_misc(OUT: str) -> None:
    with open(os.path.join(OUT, "corrupt.pdf"), "wb") as fh:
        fh.write(b"%PDF-1.7\n1 0 obj << /Type /Catalog >>\nthis is not a valid pdf body\n%%EOF")
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(os.path.join(OUT, "no_timetable.pdf"))
    c.drawString(72, 720, "Notice: the timetable will be published next week. (SYNTHETIC)")
    c.save()
    with open(os.path.join(OUT, "not_a_timetable.txt"), "w") as fh:
        fh.write("plain text")


def main(out: str | None = None) -> str:
    OUT = out or globals()["OUT"]
    os.makedirs(OUT, exist_ok=True)
    pdf_path = os.path.join(OUT, "synthetic_timetable.pdf")
    make_pdf(pdf_path)
    make_xlsx(os.path.join(OUT, "synthetic_timetable.xlsx"))
    make_csv(os.path.join(OUT, "synthetic_timetable.csv"))
    make_docx(os.path.join(OUT, "synthetic_timetable.docx"))
    make_image_and_scanned(pdf_path, os.path.join(OUT, "synthetic_page1.png"), os.path.join(OUT, "synthetic_scanned.pdf"))
    make_misc(OUT)
    return OUT


if __name__ == "__main__":
    import sys

    print("fixtures written to", main(sys.argv[1] if len(sys.argv) > 1 else None))
